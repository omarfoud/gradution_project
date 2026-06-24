import concurrent.futures
import hashlib
import io
import ipaddress
import json
import logging
import os
import re
import socket
from urllib.parse import urlparse, urljoin
import sqlite3
import threading
from contextlib import asynccontextmanager
from functools import lru_cache
from pathlib import Path
from typing import Optional, Union

import docx
import faiss
import google.generativeai as genai
import jwt
import numpy as np
import pyodbc
import PyPDF2
import requests
from cachetools.func import ttl_cache
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Header
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from sentence_transformers import SentenceTransformer

load_dotenv()

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"), format="%(asctime)s | %(levelname)s | %(message)s")
logger = logging.getLogger("job-assistant")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MODEL_NAME = os.getenv("MODEL_NAME", "all-MiniLM-L6-v2")
DEVICE = os.getenv("DEVICE", "cpu")

ARTIFACT_DIR = Path(os.getenv("HF_ARTIFACT_DIR", os.getenv("ARTIFACT_DIR", "artifacts")))
DB_PATH = os.getenv("DB_PATH", str(ARTIFACT_DIR / "jobs.db"))
INDEX_PATH = os.getenv("INDEX_PATH", str(ARTIFACT_DIR / "jobs.index"))

SQL_SERVER = os.getenv("SQL_SERVER")
SQL_PORT = os.getenv("SQL_PORT", "1433")
SQL_DATABASE = os.getenv("SQL_DATABASE")
SQL_USERNAME = os.getenv("SQL_USERNAME")
SQL_PASSWORD = os.getenv("SQL_PASSWORD")
SQL_DRIVER = os.getenv("SQL_DRIVER", "ODBC Driver 18 for SQL Server")
SQL_ENCRYPT = os.getenv("SQL_ENCRYPT", "yes")

JWT_SECRET = os.getenv("JWT_SECRET", "super-secret-key")
JWT_ALGORITHM = os.getenv("JWT_ALGORITHM", "HS256")
REQUIRE_AUTH = os.getenv("REQUIRE_AUTH", "true").lower() in ("true", "1", "yes")
FAISS_SAVE_DIRECT = os.getenv("FAISS_SAVE_DIRECT", "false").lower() in ("true", "1", "yes")
ADMIN_API_KEY = os.getenv("ADMIN_API_KEY")

INSECURE_PLACEHOLDERS = {
    "super-secret-key",
    "change-this-to-a-long-random-string",
    "change-me",
    "your_jwt_secret",
}

if not JWT_SECRET or JWT_SECRET.strip() in INSECURE_PLACEHOLDERS:
    raise RuntimeError(
        "CRITICAL SECURITY ERROR: JWT_SECRET must be set to a secure, non-placeholder value."
    )

if not ADMIN_API_KEY or ADMIN_API_KEY.strip() in INSECURE_PLACEHOLDERS:
    raise RuntimeError(
        "CRITICAL SECURITY ERROR: ADMIN_API_KEY must be set to a secure, non-placeholder value."
    )


def get_sql_identifier(env_name: str, default: str) -> str:
    value = os.getenv(env_name, default).strip()
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value):
        raise RuntimeError(f"Invalid SQL identifier in {env_name}: {value}")
    return value


APPLICANT_TABLE = get_sql_identifier("SQL_TABLE_APPLICANT", "Applicants")
RESUME_TABLE = get_sql_identifier("SQL_TABLE_RESUME", "Resumes")
COMPANY_TABLE = get_sql_identifier("SQL_TABLE_COMPANY", "Companies")
JOB_POSTING_TABLE = get_sql_identifier("SQL_TABLE_JOB_POSTING", "JobPostings")
APPLICATION_TABLE = get_sql_identifier("SQL_TABLE_APPLICATION", "Applications")
RESUME_EMBEDDING_TABLE = get_sql_identifier("SQL_TABLE_RESUME_EMBEDDING", "ResumeEmbeddings")

ALLOWED_ORIGINS = [origin.strip() for origin in os.getenv("ALLOWED_ORIGINS", "*").split(",") if origin.strip()]
ALLOWED_RESUME_DOMAINS = [d.strip().lower() for d in os.getenv("ALLOWED_RESUME_DOMAINS", "").split(",") if d.strip()]

FAISS_TOP_K = int(os.getenv("FAISS_TOP_K", "500"))
RESULTS_LIMIT = int(os.getenv("RESULTS_LIMIT", "10"))
MAX_CV_CHARS = int(os.getenv("MAX_CV_CHARS", "4000"))
REQUEST_TIMEOUT_SECONDS = int(os.getenv("REQUEST_TIMEOUT_SECONDS", "20"))
SQL_BATCH_SIZE = int(os.getenv("SQL_BATCH_SIZE", "900"))
REQUIRE_REMOTE_RESUME_URL = os.getenv("REQUIRE_REMOTE_RESUME_URL", "true").lower() in {"1", "true", "yes", "y"}
MOCK_RESUME_PATH = os.getenv("MOCK_RESUME_PATH")

MODEL: Optional[SentenceTransformer] = None
INDEX = None
INDEX_LOCK = threading.RLock()
INDEX_DIRTY = False
INDEX_AUTOSAVE_ACTIVE = True

if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)
else:
    logger.warning("GEMINI_API_KEY is not configured.")


def faiss_index_autosave_worker():
    global INDEX_DIRTY, INDEX_AUTOSAVE_ACTIVE, INDEX
    import time
    logger.info("FAISS index autosave worker started.")
    while INDEX_AUTOSAVE_ACTIVE:
        time.sleep(30)
        if INDEX_DIRTY and INDEX is not None:
            try:
                if FAISS_SAVE_DIRECT:
                    with INDEX_LOCK:
                        INDEX_DIRTY = False
                        logger.info("Saving FAISS index directly to %s...", INDEX_PATH)
                        Path(INDEX_PATH).parent.mkdir(parents=True, exist_ok=True)
                        faiss.write_index(INDEX, INDEX_PATH)
                else:
                    with INDEX_LOCK:
                        cloned_index = faiss.clone_index(INDEX)
                        INDEX_DIRTY = False
                    logger.info("Asynchronously saving FAISS index to %s...", INDEX_PATH)
                    Path(INDEX_PATH).parent.mkdir(parents=True, exist_ok=True)
                    faiss.write_index(cloned_index, INDEX_PATH)
                logger.info("FAISS index saved successfully.")
            except Exception as exc:
                logger.error("Failed to save FAISS index: %s", exc)
                INDEX_DIRTY = True
    logger.info("FAISS index autosave worker stopped.")


@asynccontextmanager
async def lifespan(app: FastAPI):
    global MODEL, INDEX, INDEX_AUTOSAVE_ACTIVE, INDEX_DIRTY
    try:
        logger.info("Loading embedding model: %s on %s", MODEL_NAME, DEVICE)
        MODEL = SentenceTransformer(MODEL_NAME, device=DEVICE)
    except Exception as exc:
        logger.error("FAILED to load embedding model '%s': %s", MODEL_NAME, exc, exc_info=True)
        MODEL = None
    if not os.path.isfile(DB_PATH) or not os.path.isfile(INDEX_PATH):
        logger.warning("jobs.db or jobs.index not found.")
    else:
        try:
            logger.info("Loading FAISS index: %s", INDEX_PATH)
            INDEX = faiss.read_index(INDEX_PATH)
            if hasattr(INDEX, "nprobe"):
                INDEX.nprobe = int(os.getenv("FAISS_NPROBE", "20"))
        except Exception as exc:
            logger.error("FAILED to load FAISS index '%s': %s", INDEX_PATH, exc, exc_info=True)
            INDEX = None

    INDEX_AUTOSAVE_ACTIVE = True
    t = threading.Thread(target=faiss_index_autosave_worker, daemon=True)
    t.start()

    logger.info("Backend is ready")
    yield

    logger.info("Shutting down backend, stopping autosave worker...")
    INDEX_AUTOSAVE_ACTIVE = False
    if INDEX_DIRTY and INDEX is not None:
        logger.info("Saving dirty FAISS index to disk before shutdown...")
        try:
            faiss.write_index(INDEX, INDEX_PATH)
            logger.info("FAISS index saved successfully on shutdown.")
        except Exception as exc:
            logger.error("Failed to save FAISS index on shutdown: %s", exc)


app = FastAPI(title="Job Assistant API", version="4.0.0", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=ALLOWED_ORIGINS, allow_methods=["*"], allow_headers=["*"])


class RecommendRequest(BaseModel):
    user_id: str = Field(..., min_length=1)
    location: Optional[str] = None
    work_type: Optional[str] = None
    experience: Optional[str] = None
    limit: int = Field(default=5, ge=1, le=50)


class AnalyzeJobRequest(BaseModel):
    user_id: str = Field(..., min_length=1)
    job_id: Union[int, str]


class SearchRequest(BaseModel):
    query: str = Field(..., min_length=1)
    location: Optional[str] = None
    work_type: Optional[str] = None
    experience: Optional[str] = None
    limit: int = Field(default=10, ge=1, le=50)


class ChatRequest(BaseModel):
    message: str = Field(..., min_length=1)
    user_id: Optional[str] = None


def check_auth(requested_id: Optional[str], auth_header: Optional[str] = None) -> Optional[dict]:
    if not auth_header:
        if REQUIRE_AUTH:
            raise HTTPException(status_code=401, detail="Authorization header is required")
        return None
    parts = auth_header.split()
    if len(parts) != 2 or parts[0].lower() != "bearer":
        raise HTTPException(status_code=401, detail="Invalid authorization header format. Use 'Bearer <token>'")
    token = parts[1]
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.PyJWTError as exc:
        raise HTTPException(status_code=401, detail=f"Invalid or expired token: {exc}") from exc
    sub = payload.get("sub")
    if not sub:
        raise HTTPException(status_code=401, detail="Token is missing 'sub' claim")
    if requested_id and str(sub) != str(requested_id):
        raise HTTPException(status_code=403, detail="Forbidden: Token subject does not match requested context")
    return payload


class CompanyContextRequest(BaseModel):
    company_user_id: Optional[str] = None
    company_id: Optional[str] = None


class CompanyChatRequest(CompanyContextRequest):
    message: str = Field(..., min_length=1)
    job_posting_id: Optional[str] = None


class CompanyJobPostRequest(CompanyContextRequest):
    job_posting_id: str


class CompanyFindCandidatesRequest(CompanyContextRequest):
    job_posting_id: str
    limit: int = Field(default=10, ge=1, le=25)


class CompanyGenerateJobPostRequest(CompanyContextRequest):
    role: str = Field(..., min_length=1)
    level: Optional[str] = None
    skills: list[str] = Field(default_factory=list)
    work_type: Optional[str] = None
    location: Optional[str] = None
    salary_range: Optional[str] = None


class SyncJobEmbeddingRequest(CompanyContextRequest):
    job_posting_id: str = Field(..., min_length=1)


class SyncResumeEmbeddingRequest(BaseModel):
    user_id: str = Field(..., min_length=1)
    force: bool = False


class CompanyHiringInsightsRequest(CompanyContextRequest):
    job_posting_id: Optional[str] = None


def get_jobs_db() -> sqlite3.Connection:
    Path(DB_PATH).parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 5000")
    return conn


def get_app_db():
    missing = [name for name, value in {"SQL_SERVER": SQL_SERVER, "SQL_DATABASE": SQL_DATABASE, "SQL_USERNAME": SQL_USERNAME, "SQL_PASSWORD": SQL_PASSWORD}.items() if not value]
    if missing:
        raise HTTPException(status_code=500, detail=f"Missing SQL Server environment variables: {', '.join(missing)}")
    conn_str = (f"DRIVER={{{SQL_DRIVER}}};SERVER={SQL_SERVER},{SQL_PORT};DATABASE={SQL_DATABASE};UID={SQL_USERNAME};PWD={SQL_PASSWORD};Encrypt={SQL_ENCRYPT};TrustServerCertificate=yes;")
    return pyodbc.connect(conn_str, timeout=10)


def get_resume_metadata_by_user_id(user_id: str) -> dict:
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        cursor.execute(f"SELECT TOP 1 a.ApplicantID, r.ResumeID, r.FilePath, r.UploadDate FROM {APPLICANT_TABLE} a JOIN {RESUME_TABLE} r ON r.ApplicantID = a.ApplicantID WHERE a.UserId = ? AND r.IsActive = 1 ORDER BY r.UploadDate DESC", user_id)
        row = cursor.fetchone()
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"SQL Server query failed: {exc}") from exc
    finally:
        conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Active resume not found for this user")
    if len(row) == 1:
        if not row[0]:
            raise HTTPException(status_code=404, detail="Active resume not found for this user")
        return {
            "user_id": str(user_id),
            "applicant_id": None,
            "resume_id": None,
            "resume_path": str(row[0]),
            "upload_date": None,
        }
    if not row[2]:
        raise HTTPException(status_code=404, detail="Active resume not found for this user")
    return {
        "user_id": str(user_id),
        "applicant_id": str(row[0]) if row[0] is not None else None,
        "resume_id": str(row[1]) if row[1] is not None else None,
        "resume_path": str(row[2]),
        "upload_date": str(row[3]) if row[3] is not None else None,
    }


def get_resume_path_by_user_id(user_id: str) -> str:
    return get_resume_metadata_by_user_id(user_id)["resume_path"]


def _validate_url_and_ips(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        raise HTTPException(status_code=400, detail="Invalid URL scheme. Only http/https are allowed.")
    hostname = parsed.hostname
    if not hostname:
        raise HTTPException(status_code=400, detail="Invalid URL (missing hostname)")

    try:
        ips = socket.getaddrinfo(hostname, None)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Failed to resolve hostname '{hostname}': {exc}")

    for ip_info in ips:
        ip_str = ip_info[4][0]
        if "%" in ip_str:
            ip_str = ip_str.split("%")[0]
        ip = ipaddress.ip_address(ip_str)
        if (
            ip.is_private or
            ip.is_loopback or
            ip.is_link_local or
            ip.is_multicast or
            ip.is_reserved or
            ip.is_unspecified
        ):
            raise HTTPException(status_code=400, detail=f"Unsafe IP address detected for {hostname} (SSRF protection block)")

    if ALLOWED_RESUME_DOMAINS:
        hostname_lower = hostname.lower()
        domain_allowed = False
        for allowed_domain in ALLOWED_RESUME_DOMAINS:
            if hostname_lower == allowed_domain or hostname_lower.endswith("." + allowed_domain):
                domain_allowed = True
                break
        if not domain_allowed:
            raise HTTPException(status_code=400, detail=f"Domain '{hostname}' is not allowed (domain allowlist block)")

    return hostname


def load_resume_bytes(resume_path: str) -> bytes:
    if resume_path.startswith(("http://", "https://")):
        url = resume_path
        max_redirects = 5
        for _ in range(max_redirects):
            _validate_url_and_ips(url)
            try:
                response = requests.get(url, timeout=REQUEST_TIMEOUT_SECONDS, allow_redirects=False)
                if response.is_redirect or response.status_code in (301, 302, 303, 307, 308):
                    location = response.headers.get("Location")
                    if not location:
                        raise HTTPException(status_code=502, detail="Redirect response missing Location header")
                    url = urljoin(url, location)
                    continue
                response.raise_for_status()
                return response.content
            except requests.RequestException as exc:
                raise HTTPException(status_code=502, detail=f"Failed to download resume: {exc}") from exc
        raise HTTPException(status_code=400, detail="Too many redirects")

    if REQUIRE_REMOTE_RESUME_URL:
        raise HTTPException(status_code=400, detail="Resume.FilePath must be an http/https URL in production.")
    if not os.path.exists(resume_path):
        raise HTTPException(status_code=404, detail=f"Resume file not found: {resume_path}")
    try:
        with open(resume_path, "rb") as file_obj:
            return file_obj.read()
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"Failed to read local resume file: {exc}") from exc


def extract_text(file_content: bytes, filename: str) -> str:
    lower_name = filename.lower().split("?")[0]
    is_pdf = lower_name.endswith(".pdf") or file_content.startswith(b"%PDF")
    
    is_docx = False
    if lower_name.endswith(".docx"):
        is_docx = True
    elif file_content.startswith(b"PK\x03\x04"):
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(file_content)) as zf:
                if "word/document.xml" in zf.namelist():
                    is_docx = True
        except Exception:
            pass

    is_html = file_content.startswith(b"<") or b"html" in file_content[:200].lower()
    try:
        if is_pdf:
            try:
                reader = PyPDF2.PdfReader(io.BytesIO(file_content))
                text = "\n".join((page.extract_text() or "") for page in reader.pages)
            except Exception as exc:
                raise HTTPException(status_code=400, detail=f"Failed to parse PDF file. Ensure it is a valid, uncorrupted PDF document. Error: {exc}")
        elif is_docx:
            try:
                document = docx.Document(io.BytesIO(file_content))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            except Exception as exc:
                raise HTTPException(status_code=400, detail=f"Failed to parse DOCX file. Ensure it is a valid, uncorrupted Word document. Error: {exc}")
        elif is_html:
            raise HTTPException(status_code=502, detail="Resume URL returned HTML, not a PDF/DOCX (possibly due to an authorization wall, redirect, or 404 error)")
        else:
            raise HTTPException(status_code=400, detail="Only PDF and DOCX resumes are supported (extension not recognized or invalid file signature)")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Failed to read resume file: {exc}") from exc
    cleaned_text = text.strip()
    if not cleaned_text:
        raise HTTPException(status_code=400, detail="Resume text extraction returned empty content")
    return cleaned_text[:MAX_CV_CHARS]


@ttl_cache(maxsize=256, ttl=600)
def get_resume_text_for_user(user_id: str) -> str:
    resume_path = MOCK_RESUME_PATH or get_resume_path_by_user_id(user_id)
    file_content = load_resume_bytes(resume_path)
    return extract_text(file_content, resume_path)


def ensure_resume_embedding_table(conn) -> None:
    cursor = conn.cursor()
    cursor.execute(
        f"""
        IF OBJECT_ID(N'dbo.{RESUME_EMBEDDING_TABLE}', N'U') IS NULL
        BEGIN
            CREATE TABLE dbo.{RESUME_EMBEDDING_TABLE} (
                UserId NVARCHAR(450) NOT NULL PRIMARY KEY,
                ApplicantID NVARCHAR(64) NULL,
                ResumeID NVARCHAR(64) NULL,
                ResumePath NVARCHAR(2048) NOT NULL,
                ModelName NVARCHAR(200) NOT NULL,
                Dimension INT NOT NULL,
                TextHash CHAR(64) NOT NULL,
                Embedding VARBINARY(MAX) NOT NULL,
                CreatedAt DATETIME2 NOT NULL CONSTRAINT DF_{RESUME_EMBEDDING_TABLE}_CreatedAt DEFAULT SYSUTCDATETIME(),
                UpdatedAt DATETIME2 NOT NULL CONSTRAINT DF_{RESUME_EMBEDDING_TABLE}_UpdatedAt DEFAULT SYSUTCDATETIME()
            );
            CREATE INDEX IX_{RESUME_EMBEDDING_TABLE}_ResumeID ON dbo.{RESUME_EMBEDDING_TABLE}(ResumeID);
            CREATE INDEX IX_{RESUME_EMBEDDING_TABLE}_TextHash ON dbo.{RESUME_EMBEDDING_TABLE}(TextHash);
        END
        """
    )
    conn.commit()


def embed_resume_text(resume_text: str) -> np.ndarray:
    if MODEL is None:
        raise HTTPException(status_code=503, detail="Embedding model is not loaded yet.")
    vector = MODEL.encode([resume_text], convert_to_numpy=True).astype("float32")
    faiss.normalize_L2(vector)
    return vector.copy()


def vector_to_bytes(vector: np.ndarray) -> bytes:
    return np.asarray(vector, dtype="float32").reshape(-1).tobytes()


def bytes_to_vector(payload: bytes, dimension: int) -> np.ndarray:
    vector = np.frombuffer(payload, dtype="float32")
    if vector.size != dimension:
        raise ValueError(f"Stored vector dimension mismatch: expected {dimension}, got {vector.size}")
    return vector.reshape(1, dimension).copy()


def sync_resume_embedding_for_user(user_id: str, *, force: bool = False) -> dict:
    metadata = get_resume_metadata_by_user_id(user_id)
    resume_path = MOCK_RESUME_PATH or metadata["resume_path"]
    file_content = load_resume_bytes(resume_path)
    resume_text = extract_text(file_content, resume_path)
    text_hash = hashlib.sha256(resume_text.encode("utf-8")).hexdigest()

    conn = get_app_db()
    try:
        ensure_resume_embedding_table(conn)
        cursor = conn.cursor()
        cursor.execute(
            f"SELECT TextHash, ModelName FROM dbo.{RESUME_EMBEDDING_TABLE} WHERE UserId = ?",
            str(user_id),
        )
        existing = cursor.fetchone()
        if existing and not force and existing[0] == text_hash and existing[1] == MODEL_NAME:
            return {
                "user_id": str(user_id),
                "resume_id": metadata["resume_id"],
                "status": "already_synced",
                "text_hash": text_hash,
                "model_name": MODEL_NAME,
            }

        vector = embed_resume_text(resume_text)
        dimension = int(vector.shape[1])
        payload = vector_to_bytes(vector)
        cursor.execute(
            f"""
            UPDATE dbo.{RESUME_EMBEDDING_TABLE}
            SET ApplicantID = ?, ResumeID = ?, ResumePath = ?, ModelName = ?,
                Dimension = ?, TextHash = ?, Embedding = ?, UpdatedAt = SYSUTCDATETIME()
            WHERE UserId = ?
            """,
            metadata["applicant_id"],
            metadata["resume_id"],
            resume_path,
            MODEL_NAME,
            dimension,
            text_hash,
            pyodbc.Binary(payload),
            str(user_id),
        )
        if cursor.rowcount == 0:
            cursor.execute(
                f"""
                INSERT INTO dbo.{RESUME_EMBEDDING_TABLE}
                    (UserId, ApplicantID, ResumeID, ResumePath, ModelName, Dimension, TextHash, Embedding)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                str(user_id),
                metadata["applicant_id"],
                metadata["resume_id"],
                resume_path,
                MODEL_NAME,
                dimension,
                text_hash,
                pyodbc.Binary(payload),
            )
        conn.commit()
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"Resume embedding sync failed: {exc}") from exc
    finally:
        conn.close()

    get_resume_text_for_user.cache_clear()
    return {
        "user_id": str(user_id),
        "applicant_id": metadata["applicant_id"],
        "resume_id": metadata["resume_id"],
        "resume_path": resume_path,
        "status": "synced",
        "text_hash": text_hash,
        "model_name": MODEL_NAME,
        "dimension": dimension,
    }


def get_stored_resume_embedding(user_id: str) -> Optional[np.ndarray]:
    conn = get_app_db()
    try:
        ensure_resume_embedding_table(conn)
        cursor = conn.cursor()
        cursor.execute(
            f"SELECT Embedding, Dimension, ModelName FROM dbo.{RESUME_EMBEDDING_TABLE} WHERE UserId = ?",
            str(user_id),
        )
        row = cursor.fetchone()
    except pyodbc.Error as exc:
        logger.warning("Failed to read stored resume embedding for user_id=%s: %s", user_id, exc)
        return None
    finally:
        conn.close()

    if not row or not row[0] or row[2] != MODEL_NAME:
        return None
    try:
        return bytes_to_vector(bytes(row[0]), int(row[1]))
    except ValueError as exc:
        logger.warning("Invalid stored resume embedding for user_id=%s: %s", user_id, exc)
        return None


def get_resume_vector_for_user(user_id: str) -> np.ndarray:
    stored_vector = get_stored_resume_embedding(user_id)
    if stored_vector is not None:
        return stored_vector
    sync_resume_embedding_for_user(user_id)
    stored_vector = get_stored_resume_embedding(user_id)
    if stored_vector is not None:
        return stored_vector
    resume_text = get_resume_text_for_user(user_id)
    return embed_resume_text(resume_text)


@lru_cache(maxsize=512)
def encode_query_cached(query_text: str):
    if MODEL is None:
        raise HTTPException(status_code=503, detail="Embedding model is not loaded yet")
    vector = MODEL.encode([query_text], convert_to_numpy=True).astype("float32")
    faiss.normalize_L2(vector)
    return vector.copy()


def fetch_jobs_by_faiss_ids(conn: sqlite3.Connection, faiss_ids: list[int], filters: dict) -> list[sqlite3.Row]:
    rows: list[sqlite3.Row] = []
    for start in range(0, len(faiss_ids), SQL_BATCH_SIZE):
        batch = faiss_ids[start: start + SQL_BATCH_SIZE]
        placeholders = ",".join(["?"] * len(batch))
        query = f"SELECT * FROM jobs WHERE faiss_id IN ({placeholders})"
        params: list = list(batch)
        conditions = []
        if filters.get("location"):
            conditions.append("location LIKE ?")
            params.append(f"%{filters['location']}%")
        if filters.get("work_type"):
            conditions.append("work_type LIKE ?")
            params.append(f"%{filters['work_type']}%")
        if filters.get("experience"):
            conditions.append("experience LIKE ?")
            params.append(f"%{filters['experience']}%")
        if conditions:
            query += " AND " + " AND ".join(conditions)
        rows.extend(conn.execute(query, params).fetchall())
    return rows


def as_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "y", "featured"}
    return False


def get_job_featured_flag(job: dict) -> bool:
    for key in ("isFeatured", "IsFeatured", "is_featured", "featured"):
        if key in job:
            return as_bool(job.get(key))
    return False


def ensure_jobs_table(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS jobs (
            faiss_id INTEGER PRIMARY KEY,
            job_id INTEGER UNIQUE,
            title TEXT,
            company TEXT,
            description TEXT,
            qualifications TEXT,
            responsibilities TEXT,
            skills TEXT,
            experience TEXT,
            location TEXT,
            work_type TEXT,
            salary_range TEXT,
            is_featured INTEGER DEFAULT 0
        )
        """
    )
    columns = {row[1] for row in conn.execute("PRAGMA table_info(jobs)").fetchall()}
    if "is_featured" not in columns:
        conn.execute("ALTER TABLE jobs ADD COLUMN is_featured INTEGER DEFAULT 0")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_job_id ON jobs(job_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_location ON jobs(location)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_work_type ON jobs(work_type)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_experience ON jobs(experience)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_is_featured ON jobs(is_featured)")
    conn.commit()


def clean_embedding_text(value: object) -> str:
    text = str(value or "")
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\\s+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def build_search_job_text(job: dict) -> str:
    return " ".join(
        clean_embedding_text(value)
        for value in [
            job.get("title"),
            job.get("company"),
            job.get("skills"),
            job.get("qualifications"),
            job.get("responsibilities"),
            job.get("description"),
            job.get("experience"),
            job.get("location"),
            job.get("work_type"),
        ]
    ).strip()


def upsert_job_embedding(job: dict) -> dict:
    global INDEX, INDEX_DIRTY
    if MODEL is None:
        raise HTTPException(status_code=503, detail="Embedding model is not loaded yet.")
    if INDEX is None:
        raise HTTPException(status_code=503, detail="FAISS index is not loaded yet.")

    job_id = job["job_id"]
    embed_text = build_search_job_text(job)
    if not embed_text:
        raise HTTPException(status_code=400, detail="Job does not contain enough text to embed.")

    with INDEX_LOCK:
        vector = MODEL.encode([embed_text], convert_to_numpy=True).astype("float32")
        faiss.normalize_L2(vector)
        faiss_id = int(getattr(INDEX, "ntotal", 0))
        INDEX.add(vector)
        INDEX_DIRTY = True

        conn = get_jobs_db()
        try:
            ensure_jobs_table(conn)
            conn.execute(
                """
                INSERT OR REPLACE INTO jobs (
                    faiss_id, job_id, title, company, description,
                    qualifications, responsibilities, skills, experience,
                    location, work_type, salary_range, is_featured
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    faiss_id,
                    job_id,
                    job.get("title"),
                    job.get("company"),
                    job.get("description"),
                    job.get("qualifications"),
                    job.get("responsibilities"),
                    job.get("skills"),
                    job.get("experience"),
                    job.get("location"),
                    job.get("work_type"),
                    job.get("salary_range"),
                    1 if get_job_featured_flag(job) else 0,
                ),
            )
            conn.commit()
        finally:
            conn.close()

        if os.getenv("TESTING") == "true":
            Path(INDEX_PATH).parent.mkdir(parents=True, exist_ok=True)
            faiss.write_index(INDEX, INDEX_PATH)

    encode_query_cached.cache_clear()
    return {"job_id": job_id, "faiss_id": faiss_id, "index_total": int(getattr(INDEX, "ntotal", faiss_id + 1))}


def search_hybrid_vector(vector: np.ndarray, *, k: int = RESULTS_LIMIT, location: Optional[str] = None, work_type: Optional[str] = None, experience: Optional[str] = None):
    if MODEL is None or INDEX is None:
        raise HTTPException(status_code=503, detail="FAISS search is not loaded yet.")
    if not os.path.exists(DB_PATH):
        raise HTTPException(status_code=503, detail="jobs.db is missing.")

    search_k = FAISS_TOP_K
    if location or work_type or experience:
        search_k = max(search_k, 2000)
    ntotal = int(getattr(INDEX, "ntotal", 0))
    if ntotal > 0:
        search_k = min(search_k, ntotal)

    with INDEX_LOCK:
        distances, indices = INDEX.search(vector, search_k)
        
    valid_ids = [int(idx) for idx in indices[0] if idx != -1]
    if not valid_ids:
        return []
    score_by_id = {int(idx): float(score) for idx, score in zip(indices[0], distances[0]) if idx != -1}
    id_to_rank = {idx: rank for rank, idx in enumerate(valid_ids)}
    conn = get_jobs_db()
    try:
        rows = fetch_jobs_by_faiss_ids(conn, valid_ids, {"location": location, "work_type": work_type, "experience": experience})
    finally:
        conn.close()
    sorted_rows = sorted(
        rows,
        key=lambda r: (
            0 if as_bool(r["is_featured"]) else 1,
            id_to_rank.get(r["faiss_id"], 10**9),
        ),
    )
    results = []
    seen = set()
    for row in sorted_rows:
        job_id = row["job_id"]
        if job_id in seen:
            continue
        item = dict(row)
        item["similarity_score"] = round(score_by_id.get(row["faiss_id"], 0.0), 4)
        item["embedding_source"] = "faiss_local"
        item["isFeatured"] = as_bool(item.get("is_featured"))
        results.append(item)
        seen.add(job_id)
        if len(results) >= k:
            break
    return results


def search_hybrid(query_text: str, *, k: int = RESULTS_LIMIT, location: Optional[str] = None, work_type: Optional[str] = None, experience: Optional[str] = None):
    vector = encode_query_cached(query_text.strip())
    return search_hybrid_vector(vector, k=k, location=location, work_type=work_type, experience=experience)


def recommend_jobs_for_user(user_id: str, *, k: int = RESULTS_LIMIT, location: Optional[str] = None, work_type: Optional[str] = None, experience: Optional[str] = None):
    vector = get_resume_vector_for_user(user_id)
    return search_hybrid_vector(vector, k=k, location=location, work_type=work_type, experience=experience)


def build_analysis_prompt(job_row: sqlite3.Row, resume_text: str) -> str:
    job_details = f"""
- Title: {job_row['title']}
- Company: {job_row['company']}
- Location: {job_row['location']}
- Work Type: {job_row['work_type']}
- Experience: {job_row['experience']}
- Salary Range: {job_row['salary_range']}
- Skills: {job_row['skills']}
- Qualifications: {job_row['qualifications']}
- Responsibilities: {job_row['responsibilities']}
- Description: {job_row['description']}
""".strip()
    return f"You are an expert career-matching assistant.\n\nJob details:\n{job_details}\n\nCandidate CV:\n{resume_text[:2200]}\n\nReturn:\n1. Match %\n2. Why\n3. Strengths\n4. Gaps\n5. Plan".strip()


def generate_text(prompt: str) -> str:
    if not GEMINI_API_KEY:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY is not configured")
    try:
        model = genai.GenerativeModel(GEMINI_MODEL)
        response = model.generate_content(prompt)
        return response.text or "No response generated."
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"LLM request failed: {exc}") from exc


def generate_text_or_fallback(prompt: str, fallback: str) -> str:
    try:
        return generate_text(prompt)
    except HTTPException as exc:
        if exc.status_code == 502:
            logger.warning("Using fallback text because LLM request failed: %s", exc.detail)
            return fallback
        raise


def build_recommendation_reply_fallback(jobs: list[dict]) -> str:
    if not jobs:
        return "I could not find matching jobs right now."
    lines = ["Here are the top matching jobs from your CV embedding:"]
    for idx, job in enumerate(jobs[:5], 1):
        score = job.get("similarity_score")
        score_text = f" ({round(float(score) * 100)}% similarity)" if score is not None else ""
        lines.append(f"{idx}. {job.get('title')} at {job.get('company')}{score_text}")
    return "\n".join(lines)


def build_analysis_fallback(job_row: sqlite3.Row, resume_text: str) -> str:
    score = calculate_cv_match_score(build_job_match_text(dict(job_row)), resume_text)
    return (
        f"Match Score: {score}%\n"
        f"Why: This fallback score is based on semantic similarity between the CV and the job text.\n"
        f"Strengths: Skills and terms that overlap with the job requirements are counted positively.\n"
        f"Gaps: Review the job requirements manually for missing tools, seniority, or domain-specific experience.\n"
        f"Plan: Improve the CV by highlighting directly relevant projects, skills, and measurable outcomes for this role."
    )


def build_candidate_ranking_fallback(candidates: list[dict]) -> str:
    if not candidates:
        return "No applicants found for this job posting."
    lines = ["AI candidate ranking fallback based on computed CV/job match scores:"]
    for idx, candidate in enumerate(candidates[:10], 1):
        lines.append(f"{idx}. {candidate.get('name') or 'Candidate'} - {candidate.get('match_score', 0)}% match")
    return "\n".join(lines)


def ensure_company_context(company_user_id: Optional[str], company_id: Optional[str], sub: Optional[str] = None) -> dict:
    if not company_user_id and company_id is None and not sub:
        raise HTTPException(status_code=400, detail="Send company_user_id from logged-in company session or company_id for local testing")
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        if company_id is not None:
            cursor.execute(f"SELECT TOP 1 CompanyID, Name, Industry, WebsiteURL, HeadquarterAddress, Location, LogoUrl, UserId FROM {COMPANY_TABLE} WHERE CompanyID = ?", str(company_id))
        elif company_user_id is not None:
            cursor.execute(f"SELECT TOP 1 CompanyID, Name, Industry, WebsiteURL, HeadquarterAddress, Location, LogoUrl, UserId FROM {COMPANY_TABLE} WHERE UserId = ?", company_user_id)
        else:
            cursor.execute(f"SELECT TOP 1 CompanyID, Name, Industry, WebsiteURL, HeadquarterAddress, Location, LogoUrl, UserId FROM {COMPANY_TABLE} WHERE UserId = ?", sub)
        row = cursor.fetchone()
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"Company lookup failed: {exc}") from exc
    finally:
        conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Company not found")
    company = {"company_id": str(row[0]), "name": row[1], "industry": row[2], "website_url": row[3], "headquarter_address": row[4], "location": row[5], "logo_url": row[6], "user_id": row[7]}
    
    if sub is not None:
        if str(company["user_id"]) != str(sub):
            raise HTTPException(status_code=403, detail="Forbidden: You do not have access to this company context")
            
    return company


def get_company_job_posting(company_id: str, job_posting_id: str) -> dict:
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        cursor.execute(f"SELECT TOP 1 JobID, Title, Description, Responsibility, MinSalary, MaxSalary, PostedDate, IsActive, IsRemote, CompanyID, JobTypes FROM {JOB_POSTING_TABLE} WHERE JobID = ? AND CompanyID = ?", str(job_posting_id), str(company_id))
        row = cursor.fetchone()
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"Job posting lookup failed: {exc}") from exc
    finally:
        conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Job posting not found for this company")
    salary_range = "Negotiable"
    if row[4] is not None and row[5] is not None:
        salary_range = f"{float(row[4]):.2f} - {float(row[5]):.2f}"
    return {"job_id": str(row[0]), "title": row[1], "description": row[2], "responsibilities": row[3], "salary_range": salary_range, "posted_date": str(row[6]) if row[6] else None, "is_active": bool(row[7]), "is_remote": bool(row[8]), "company_id": str(row[9]), "job_type": str(row[10]) if row[10] is not None else None}


def get_company_summary_for_prompt(company: dict) -> str:
    return f"Company:\n- ID: {company['company_id']}\n- Name: {company['name']}\n- Industry: {company['industry']}".strip()


def get_company_kpis(company_id: str, job_posting_id: Optional[str] = None) -> dict:
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        
        q_count = f"SELECT COUNT(*) FROM {APPLICATION_TABLE} a JOIN {JOB_POSTING_TABLE} jp ON jp.JobID = a.JobPostingID WHERE jp.CompanyID = ?"
        params = [str(company_id)]
        if job_posting_id:
            q_count += " AND jp.JobID = ?"
            params.append(str(job_posting_id))
        cursor.execute(q_count, params)
        applications_count = int(cursor.fetchone()[0])
        
        q_status = f"SELECT a.ApplicationStatus, COUNT(*) FROM {APPLICATION_TABLE} a JOIN {JOB_POSTING_TABLE} jp ON jp.JobID = a.JobPostingID WHERE jp.CompanyID = ?"
        params_status = [str(company_id)]
        if job_posting_id:
            q_status += " AND jp.JobID = ?"
            params_status.append(str(job_posting_id))
        q_status += " GROUP BY a.ApplicationStatus"
        cursor.execute(q_status, params_status)
        status_rows = cursor.fetchall()
        status_breakdown = [{"status": str(r[0]) if r[0] is not None else "Unknown", "count": int(r[1])} for r in status_rows]
        
        q_top = f"SELECT TOP 5 jp.JobID, jp.Title, COUNT(*) AS app_count FROM {APPLICATION_TABLE} a JOIN {JOB_POSTING_TABLE} jp ON jp.JobID = a.JobPostingID WHERE jp.CompanyID = ? GROUP BY jp.JobID, jp.Title ORDER BY app_count DESC"
        cursor.execute(q_top, str(company_id))
        top_rows = cursor.fetchall()
        top_jobs = [{"job_id": str(r[0]), "title": str(r[1]), "applications_count": int(r[2])} for r in top_rows]
        
        return {
            "applications_count": applications_count,
            "status_breakdown": status_breakdown,
            "top_jobs": top_jobs
        }
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"Company KPI query failed: {exc}") from exc
    finally:
        conn.close()


def sql_column_exists(cursor, table_name: str, column_name: str) -> bool:
    cursor.execute("SELECT COL_LENGTH(?, ?)", table_name, column_name)
    return cursor.fetchone()[0] is not None


def fetch_sql_job_for_embedding(job_posting_id: str, company_id: Optional[str] = None) -> dict:
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        job_location_expr = "jp.Location" if sql_column_exists(cursor, JOB_POSTING_TABLE, "Location") else "c.Location"
        qualifications_expr = "jp.Qualifications" if sql_column_exists(cursor, JOB_POSTING_TABLE, "Qualifications") else "CAST(NULL AS NVARCHAR(MAX))"
        skills_expr = "jp.Skills" if sql_column_exists(cursor, JOB_POSTING_TABLE, "Skills") else "CAST(NULL AS NVARCHAR(MAX))"
        experience_expr = "jp.Experience" if sql_column_exists(cursor, JOB_POSTING_TABLE, "Experience") else "CAST(NULL AS NVARCHAR(MAX))"
        is_featured_expr = "jp.IsFeatured" if sql_column_exists(cursor, JOB_POSTING_TABLE, "IsFeatured") else "CAST(0 AS BIT)"
        where = "jp.JobID = ?"
        params = [str(job_posting_id)]
        if company_id is not None:
            where += " AND jp.CompanyID = ?"
            params.append(str(company_id))
        cursor.execute(
            f"""
            SELECT TOP 1
                jp.JobID,
                jp.Title,
                c.Name,
                jp.Description,
                {qualifications_expr} AS Qualifications,
                jp.Responsibility,
                {skills_expr} AS Skills,
                {experience_expr} AS Experience,
                {job_location_expr} AS Location,
                jp.JobTypes,
                jp.MinSalary,
                jp.MaxSalary,
                jp.IsRemote,
                {is_featured_expr} AS IsFeatured
            FROM {JOB_POSTING_TABLE} jp
            LEFT JOIN {COMPANY_TABLE} c ON c.CompanyID = jp.CompanyID
            WHERE {where}
            """,
            *params,
        )
        row = cursor.fetchone()
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"Job embedding lookup failed: {exc}") from exc
    finally:
        conn.close()

    if not row:
        raise HTTPException(status_code=404, detail="Job posting not found")

    salary_range = "Negotiable"
    if row[10] is not None and row[11] is not None:
        salary_range = f"{float(row[10]):.2f} - {float(row[11]):.2f}"
    work_type = str(row[9]) if row[9] is not None else ("Remote" if row[12] else None)
    return {
        "job_id": str(row[0]),
        "title": row[1],
        "company": row[2],
        "description": row[3],
        "qualifications": row[4],
        "responsibilities": row[5],
        "skills": row[6],
        "experience": row[7],
        "location": row[8],
        "work_type": work_type,
        "salary_range": salary_range,
        "isFeatured": as_bool(row[13]),
    }


def get_candidates_for_job(company_id: str, job_posting_id: str, limit: int) -> list[dict]:
    limit = max(1, min(int(limit), 25))
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        match_score_expr = "app.MatchScore" if sql_column_exists(cursor, APPLICATION_TABLE, "MatchScore") else "CAST(NULL AS INT)"
        cursor.execute(f"SELECT TOP (?) app.ApplicationID, applicant.ApplicantID, applicant.FirstName, applicant.LastName, applicant.Location, app.ApplicationStatus, app.AppliedDate, COALESCE(submittedResume.FilePath, activeResume.FilePath) AS ResumePath, {match_score_expr} AS MatchScore FROM {APPLICATION_TABLE} app JOIN {JOB_POSTING_TABLE} jp ON jp.JobID = app.JobPostingID JOIN {APPLICANT_TABLE} applicant ON applicant.ApplicantID = app.ApplicantID LEFT JOIN {RESUME_TABLE} submittedResume ON submittedResume.ResumeID = app.ResumeID LEFT JOIN {RESUME_TABLE} activeResume ON activeResume.ApplicantID = applicant.ApplicantID AND activeResume.IsActive = 1 WHERE jp.CompanyID = ? AND jp.JobID = ? ORDER BY app.AppliedDate DESC", limit, str(company_id), str(job_posting_id))
        rows = cursor.fetchall()
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"Candidates query failed: {exc}") from exc
    finally:
        conn.close()
    return [{"application_id": r[0], "applicant_id": r[1], "name": f"{r[2] or ''} {r[3] or ''}".strip(), "location": r[4], "application_status": str(r[5]) if r[5] is not None else None, "applied_date": str(r[6]) if r[6] else None, "resume_path": r[7], "match_score": int(r[8]) if r[8] is not None else None} for r in rows]


def build_job_match_text(job: dict) -> str:
    return " ".join(clean_embedding_text(value) for value in [job.get("title"), job.get("description"), job.get("responsibilities"), job.get("job_type"), job.get("salary_range")]).strip()


def keyword_match_score(job_text: str, resume_text: str) -> int:
    job_terms = {term.lower() for term in re.findall(r"[A-Za-z][A-Za-z+#.\-]{1,}|[\u0600-\u06FF]{2,}", job_text) if len(term) >= 2}
    resume_terms = {term.lower() for term in re.findall(r"[A-Za-z][A-Za-z+#.\-]{1,}|[\u0600-\u06FF]{2,}", resume_text) if len(term) >= 2}
    if not job_terms or not resume_terms:
        return 0
    overlap = len(job_terms & resume_terms)
    coverage = overlap / len(job_terms)
    precision = overlap / len(resume_terms)
    score = (coverage * 0.8) + (precision * 0.2)
    return max(0, min(100, round(score * 100)))


def calculate_cv_match_score(job_text: str, resume_text: str) -> int:
    job_text = (job_text or "").strip()
    resume_text = (resume_text or "").strip()
    if not job_text or not resume_text or resume_text.startswith(("Could not read resume:", "Failed to fetch resume:")):
        return 0
    if MODEL is None:
        return keyword_match_score(job_text, resume_text)
    try:
        embeddings = MODEL.encode([job_text, resume_text], convert_to_numpy=True).astype("float32")
        faiss.normalize_L2(embeddings)
        similarity = float(embeddings[0] @ embeddings[1])
        return max(0, min(100, round(similarity * 100)))
    except Exception as exc:
        logger.warning("Embedding match score failed, using keyword fallback: %s", exc)
        return keyword_match_score(job_text, resume_text)


def add_resume_previews_and_match_scores(job: dict, candidates: list[dict]) -> list[dict]:
    job_text = build_job_match_text(job)
    def enrich_candidate(candidate: dict) -> dict:
        resume_preview = ""
        path = candidate.get("resume_path")
        if path:
            try:
                resume_preview = extract_text(load_resume_bytes(str(path)), str(path))[:2200]
            except Exception as exc:
                resume_preview = f"Could not read resume: {exc}"
        computed_score = calculate_cv_match_score(job_text, resume_preview)
        return {**candidate, "resume_preview": resume_preview, "match_score": computed_score}
    max_workers = min(len(candidates), 10) if candidates else 1
    import math
    num_batches = math.ceil(len(candidates) / max_workers) if candidates else 1
    timeout_seconds = max(20.0, num_batches * 15.0)

    enriched_candidates = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [executor.submit(enrich_candidate, candidate) for candidate in candidates]
        concurrent.futures.wait(futures, timeout=timeout_seconds)
        for future, candidate in zip(futures, candidates):
            if future.done():
                try:
                    enriched_candidates.append(future.result())
                except Exception as exc:
                    enriched_candidates.append({**candidate, "resume_preview": f"Failed to fetch resume: {exc}", "match_score": 0})
            else:
                future.cancel()
                enriched_candidates.append({**candidate, "resume_preview": "Failed to fetch resume: Request timed out", "match_score": 0})
    return sorted(enriched_candidates, key=lambda item: item.get("match_score") or 0, reverse=True)


def save_candidate_match_scores(candidates: list[dict]) -> None:
    if not candidates:
        return
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        if not sql_column_exists(cursor, APPLICATION_TABLE, "MatchScore"):
            return
        for candidate in candidates:
            cursor.execute(f"UPDATE {APPLICATION_TABLE} SET MatchScore = ? WHERE ApplicationID = ?", int(candidate.get("match_score") or 0), candidate.get("application_id"))
        conn.commit()
    except pyodbc.Error as exc:
        logger.warning("Failed to save candidate MatchScore values: %s", exc)
    finally:
        conn.close()


def build_company_chat_prompt(company: dict, message: str, job: Optional[dict] = None, kpis: Optional[dict] = None) -> str:
    job_part = f"\nRelated job posting:\n{job}" if job else ""
    kpi_part = f"\nCompany KPIs:\n{kpis}" if kpis else ""
    return f"You are an AI hiring assistant.\n{get_company_summary_for_prompt(company)}{job_part}{kpi_part}\nRecruiter question:\n{message}".strip()


def build_job_post_analysis_prompt(company: dict, job: dict) -> str:
    return f"You are an expert recruitment copywriter.\n{get_company_summary_for_prompt(company)}\nJob posting:\n{job}\nReturn: 1. Score 2. Clarity 3. Missing 4. Bias 5. Improved desc 6. Action list".strip()


def build_candidates_prompt(company: dict, job: dict, candidates: list[dict]) -> str:
    return f"You are an AI recruiter assistant.\n{get_company_summary_for_prompt(company)}\nJob: {job}\nApplicants: {candidates}\nReturn ranked_candidates with match_score, strengths, gaps, recommendation.".strip()


def summarize_recommended_jobs(jobs: list[dict]) -> list[dict]:
    return [{
        "job_id": j["job_id"],
        "title": j["title"],
        "company": j["company"],
        "location": j["location"],
        "work_type": j["work_type"],
        "experience": j["experience"],
        "salary_range": j["salary_range"],
        "similarity_score": j["similarity_score"],
        "isFeatured": as_bool(j.get("isFeatured", j.get("is_featured"))),
    } for j in jobs]


def is_recommendation_message(message: str) -> bool:
    normalized = message.lower()
    english_phrases = [
        "recommend",
        "recommendation",
        "recommended for you",
        "suggest jobs",
        "job suggestions",
        "job matches",
        "matching jobs",
        "find jobs",
        "show me jobs",
        "show me remote jobs",
        "match my skills",
        "match my cv",
        "match my resume",
        "suitable jobs",
        "jobs for me",
        "best jobs",
    ]
    arabic_phrases = [
        "رشح",
        "ترشح",
        "ترشيح",
        "وظائف مناسبة",
        "وظايف مناسبة",
        "شغل مناسب",
        "فرص مناسبة",
        "انسب وظائف",
        "أفضل وظائف",
        "افضل وظائف",
    ]
    return any(phrase in normalized for phrase in english_phrases + arabic_phrases)


def is_greeting_message(message: str) -> bool:
    normalized = re.sub(r"[^\w\s\u0600-\u06ff]", "", message.lower()).strip()
    greetings = {
        "hi",
        "hello",
        "hey",
        "good morning",
        "good afternoon",
        "good evening",
        "السلام عليكم",
        "سلام",
        "اهلا",
        "اهلاً",
        "هاي",
    }
    return normalized in greetings or len(normalized.split()) <= 2 and normalized in greetings


def require_admin_key(x_admin_key: Optional[str]) -> None:
    if ADMIN_API_KEY:
        if not x_admin_key or x_admin_key != ADMIN_API_KEY:
            raise HTTPException(status_code=401, detail="Unauthorized: Invalid or missing X-Admin-API-Key header")


@app.get("/health")
def health_check():
    return {"status": "ok", "model_loaded": MODEL is not None, "index_loaded": INDEX is not None, "jobs_db_exists": os.path.exists(DB_PATH), "index_exists": os.path.exists(INDEX_PATH), "sql_server_configured": all([SQL_SERVER, SQL_DATABASE, SQL_USERNAME, SQL_PASSWORD]), "gemini_configured": bool(GEMINI_API_KEY), "require_remote_resume_url": REQUIRE_REMOTE_RESUME_URL}


@app.get("/health/app-db")
def app_db_health_check():
    conn = get_app_db()
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT 1")
        value = cursor.fetchone()[0]
        return {"status": "ok", "sql_server_connection": bool(value)}
    except pyodbc.Error as exc:
        raise HTTPException(status_code=500, detail=f"SQL Server connection failed: {exc}") from exc
    finally:
        conn.close()


@app.post("/recommend-matches")
def recommend_matches(req: RecommendRequest, authorization: Optional[str] = Header(None)):
    check_auth(req.user_id, authorization)
    jobs = recommend_jobs_for_user(req.user_id, k=req.limit, location=req.location, work_type=req.work_type, experience=req.experience)
    return {"jobs": summarize_recommended_jobs(jobs)}


@app.post("/analyze-job-id")
def analyze_job_id(req: AnalyzeJobRequest, authorization: Optional[str] = Header(None)):
    check_auth(req.user_id, authorization)
    resume_text = get_resume_text_for_user(req.user_id)
    conn = get_jobs_db()
    try:
        row = conn.execute("SELECT * FROM jobs WHERE job_id = ?", (req.job_id,)).fetchone()
    finally:
        conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Job ID not found")
    return {"analysis": generate_text_or_fallback(build_analysis_prompt(row, resume_text), build_analysis_fallback(row, resume_text))}


@app.post("/search")
def search_jobs(req: SearchRequest, authorization: Optional[str] = Header(None)):
    check_auth(None, authorization)
    jobs = search_hybrid(req.query, k=req.limit, location=req.location, work_type=req.work_type, experience=req.experience)
    return {"jobs": jobs}


@app.post("/chat")
def chat_general(req: ChatRequest, authorization: Optional[str] = Header(None)):
    resume_context = ""
    recommendation_context = ""
    recommended_jobs: list[dict] = []
    wants_recommendations = is_recommendation_message(req.message)
    if req.user_id:
        check_auth(req.user_id, authorization)
    if is_greeting_message(req.message):
        return {
            "reply": (
                "Hello! I can help with CV improvement, interview preparation, "
                "career questions, and job recommendations based on your uploaded CV. "
                "Ask me what you want to do next."
            )
        }
    if req.user_id:
        try:
            if wants_recommendations:
                jobs = recommend_jobs_for_user(req.user_id, k=5)
                recommended_jobs = summarize_recommended_jobs(jobs)
                recommendation_context = "\nRecommended jobs from the vector search result:\n" + json.dumps(recommended_jobs, ensure_ascii=False, indent=2) + "\n"
            else:
                resume_text = get_resume_text_for_user(req.user_id)
                resume_context = f"\nCandidate CV context:\n{resume_text[:2200]}\n"
        except Exception as exc:
            logger.warning("Could not load resume for user_id=%s: %s", req.user_id, exc)
    system_prompt = (
        "You are a helpful career coach. Answer clearly and briefly. "
        "Use clean plain text with short sections and bullet points when useful."
        " If recommended jobs are provided, use those exact jobs and mention why they fit the candidate."
        f"{resume_context}{recommendation_context}\nUser message: {req.message}"
    )
    fallback_reply = build_recommendation_reply_fallback(recommended_jobs) if recommended_jobs else "AI text generation is temporarily unavailable. Please try again later."
    reply = generate_text_or_fallback(system_prompt, fallback_reply)
    response = {"reply": reply}
    if recommended_jobs:
        response["recommended_jobs"] = recommended_jobs
    return response


@app.post("/chat/company/message")
def company_chat(req: CompanyChatRequest, authorization: Optional[str] = Header(None)):
    payload = check_auth(req.company_user_id, authorization)
    sub = payload.get("sub") if payload else None
    company = ensure_company_context(req.company_user_id, req.company_id, sub=sub)
    job = get_company_job_posting(company["company_id"], req.job_posting_id) if req.job_posting_id else None
    kpis = get_company_kpis(company["company_id"], req.job_posting_id)
    reply = generate_text_or_fallback(
        build_company_chat_prompt(company, req.message, job=job, kpis=kpis),
        "AI text generation is temporarily unavailable, but the company and job context were loaded successfully.",
    )
    return {"reply": reply, "company": company, "job": job, "kpis": kpis}


@app.post("/chat/company/analyze-company")
def analyze_company(req: CompanyContextRequest, authorization: Optional[str] = Header(None)):
    payload = check_auth(req.company_user_id, authorization)
    sub = payload.get("sub") if payload else None
    company = ensure_company_context(req.company_user_id, req.company_id, sub=sub)
    kpis = get_company_kpis(company["company_id"])
    prompt = build_company_chat_prompt(company, "Analyze our company profile for hiring attractiveness.", kpis=kpis)
    return {"analysis": generate_text_or_fallback(prompt, "Company context and hiring KPIs loaded successfully. LLM analysis is temporarily unavailable."), "company": company, "kpis": kpis}


@app.post("/chat/company/analyze-job-post")
def analyze_company_job_post(req: CompanyJobPostRequest, authorization: Optional[str] = Header(None)):
    payload = check_auth(req.company_user_id, authorization)
    sub = payload.get("sub") if payload else None
    company = ensure_company_context(req.company_user_id, req.company_id, sub=sub)
    job = get_company_job_posting(company["company_id"], req.job_posting_id)
    return {"analysis": generate_text_or_fallback(build_job_post_analysis_prompt(company, job), "Job post context loaded successfully. LLM job-post analysis is temporarily unavailable."), "company": company, "job": job}


@app.post("/chat/company/hiring-insights")
def company_hiring_insights(req: CompanyHiringInsightsRequest, authorization: Optional[str] = Header(None)):
    payload = check_auth(req.company_user_id, authorization)
    sub = payload.get("sub") if payload else None
    company = ensure_company_context(req.company_user_id, req.company_id, sub=sub)
    kpis = get_company_kpis(company["company_id"], req.job_posting_id)
    prompt = build_company_chat_prompt(company, "Turn these hiring KPIs into useful recruiter insights.", kpis=kpis)
    return {"insights": generate_text_or_fallback(prompt, "Hiring KPI context loaded successfully. LLM insights are temporarily unavailable."), "company": company, "kpis": kpis}


@app.post("/chat/company/generate-job-post")
def generate_company_job_post(req: CompanyGenerateJobPostRequest, authorization: Optional[str] = Header(None)):
    payload = check_auth(req.company_user_id, authorization)
    sub = payload.get("sub") if payload else None
    company = ensure_company_context(req.company_user_id, req.company_id, sub=sub)
    prompt = f"Create a professional job post for {company['name']}.\nRole: {req.role}\nSkills: {req.skills}"
    return {"job_post": generate_text_or_fallback(prompt, f"Draft job post fallback for {req.role}: describe responsibilities, requirements, skills ({', '.join(req.skills)}), work type, location, and salary range."), "company": company}


@app.post("/chat/company/find-candidates")
def company_find_candidates(req: CompanyFindCandidatesRequest, authorization: Optional[str] = Header(None)):
    payload = check_auth(req.company_user_id, authorization)
    sub = payload.get("sub") if payload else None
    company = ensure_company_context(req.company_user_id, req.company_id, sub=sub)
    job = get_company_job_posting(company["company_id"], req.job_posting_id)
    candidates = get_candidates_for_job(company["company_id"], req.job_posting_id, req.limit)
    if not candidates:
        return {"ranking": "No applicants found for this job posting.", "company": company, "job": job, "candidates": []}
    scored_candidates = add_resume_previews_and_match_scores(job, candidates)
    save_candidate_match_scores(scored_candidates)
    ranking = generate_text_or_fallback(build_candidates_prompt(company, job, scored_candidates), build_candidate_ranking_fallback(scored_candidates))
    public_candidates = [{k: v for k, v in c.items() if k not in {"resume_path", "resume_preview"}} for c in scored_candidates]
    return {"ranking": ranking, "company": company, "job": job, "candidates": public_candidates}


@app.post("/admin/sync-job-embedding")
def sync_job_embedding(req: SyncJobEmbeddingRequest, x_admin_key: Optional[str] = Header(None, alias="X-Admin-API-Key")):
    require_admin_key(x_admin_key)
    company_id = req.company_id
    if req.company_user_id and not company_id:
        company_id = ensure_company_context(req.company_user_id, None)["company_id"]
    job = fetch_sql_job_for_embedding(req.job_posting_id, company_id)
    result = upsert_job_embedding(job)
    return {"status": "synced", "job": job, **result}


@app.post("/admin/sync-resume-embedding")
def sync_resume_embedding(req: SyncResumeEmbeddingRequest, x_admin_key: Optional[str] = Header(None, alias="X-Admin-API-Key")):
    require_admin_key(x_admin_key)
    result = sync_resume_embedding_for_user(req.user_id, force=req.force)
    return result


@app.post("/clear-cache")
def clear_cache(x_admin_key: Optional[str] = Header(None, alias="X-Admin-API-Key")):
    require_admin_key(x_admin_key)
    get_resume_text_for_user.cache_clear()
    encode_query_cached.cache_clear()
    return {"status": "cache cleared"}


@app.post("/admin/invalidate-resume-cache")
def invalidate_resume_cache(user_id: str, x_admin_key: Optional[str] = Header(None, alias="X-Admin-API-Key")):
    require_admin_key(x_admin_key)
    try:
        get_resume_text_for_user.cache.pop((user_id,), None)
        conn = get_app_db()
        try:
            ensure_resume_embedding_table(conn)
            conn.cursor().execute(f"DELETE FROM dbo.{RESUME_EMBEDDING_TABLE} WHERE UserId = ?", str(user_id))
            conn.commit()
        finally:
            conn.close()
    except Exception as exc:
        logger.warning("Failed to invalidate resume cache for user_id=%s: %s", user_id, exc)
        raise HTTPException(status_code=500, detail=f"Failed to invalidate cache: {exc}")
    return {"status": "cache invalidated", "resume_embedding": "deleted", "user_id": user_id}
